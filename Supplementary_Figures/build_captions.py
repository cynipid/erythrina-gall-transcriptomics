"""Rebuild all caption text files and the Word caption document from captions.json."""
from __future__ import annotations

import json
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent

class CaptionParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.depth={'i':0,'b':0,'sub':0,'sup':0}
        self.parts=[]
    def handle_starttag(self,tag,attrs):
        if tag in self.depth:self.depth[tag]+=1
    def handle_endtag(self,tag):
        if tag in self.depth:self.depth[tag]=max(0,self.depth[tag]-1)
    def handle_data(self,text):
        self.parts.append((text,{key:bool(value)for key,value in self.depth.items()}))

def parse(text):
    p=CaptionParser();p.feed(text);p.close();return p.parts

def plain(parts):
    out=[]
    subs=str.maketrans('0123456789','₀₁₂₃₄₅₆₇₈₉')
    for text,style in parts:
        out.append(text.translate(subs)if style['sub']else text)
    return ''.join(out)

def main():
    data=json.loads((ROOT/'captions.json').read_text())
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.85)
    for name in ('Normal','Caption','Title'):
        style=doc.styles[name]
        style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(14 if name=='Title'else 11)
        style.font.italic=False
        style.font.bold=name=='Title'
        fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
        for attr in ('ascii','hAnsi','eastAsia','cs'):
            fonts.set(qn('w:'+attr),'Times New Roman')
        ppr=style.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Caption'].paragraph_format.line_spacing=1.15
    doc.styles['Caption'].paragraph_format.space_after=Pt(8)
    doc.styles['Caption'].paragraph_format.widow_control=True
    doc.styles['Caption'].paragraph_format.keep_together=True
    title=doc.add_paragraph('Supplementary figure captions',style='Title')
    title.runs[0].font.name='Times New Roman'
    title.runs[0].font.size=Pt(14)
    combined=[]
    for figure in data['figures']:
        parts=parse(figure['caption_html'])
        text=plain(parts)
        assert text==figure['caption'],figure['label']+' plain/HTML mismatch'
        (ROOT/figure['folder']/'caption.txt').write_text(text+'\n',encoding='utf-8')
        combined.append(text)
        para=doc.add_paragraph(style='Caption')
        for content,style in parts:
            run=para.add_run(content)
            run.italic=style['i'];run.bold=style['b']
            run.font.subscript=style['sub'];run.font.superscript=style['sup']
    (ROOT/'Supplementary_Figure_Captions.txt').write_text('\n\n'.join(combined)+'\n',encoding='utf-8')
    doc.core_properties.title='Supplementary figure captions'
    doc.core_properties.author=data['author']
    doc.save(ROOT/'Supplementary_Figure_Captions.docx')
    print('Rebuilt eight individual captions, combined TXT and DOCX')

if __name__=='__main__':main()
